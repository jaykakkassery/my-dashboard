#!/usr/bin/env python3
"""
Hotel Booked Property Information Report
Compares hotel name/address/phone from HOTEL_BOOKING_V2 vs supplier reference data (MongoDB).

Usage:
  python3 hotel_address_report.py [DDMMYYYY]

  DDMMYYYY  — date to report on (e.g. 02092026 for 2 Sep 2026).
              Defaults to today if omitted.
"""

import os
import pymysql
import pymongo
import re
import sys
import unicodedata
from difflib import SequenceMatcher
from datetime import datetime

# ---------- date handling ----------
# Accept DDMMYYYY (e.g. 02092026); fall back to today.
_arg = sys.argv[1] if len(sys.argv) > 1 else None
if _arg:
    try:
        _dt = datetime.strptime(_arg, '%d%m%Y')
    except ValueError:
        print(f"ERROR: expected DDMMYYYY, got {_arg!r}")
        sys.exit(1)
else:
    _dt = datetime.now()

DDMMYYYY    = _dt.strftime('%d%m%Y')      # for the filename
REPORT_DATE = _dt.strftime('%Y-%m-%d')    # for SQL / display
DISPLAY_DATE = _dt.strftime('%d %b %Y')   # human-readable in the report heading

MYSQL_HOST     = 'mysqlread.prod.infra.tstllc.net'
MYSQL_PORT     = 3306
MYSQL_USER     = 'v-ldap-jayap-aws-dev-0NbBxGbN0C6'
MYSQL_PASSWORD = 'zf1zbnoY2fyy-EbLpnvD'
MYSQL_DB       = 'book'
MONGO_URI      = 'mongodb://infra1-internal.infra.tstllc.net:31010/'

GDRIVE_FOLDER_ID  = '1U2Xz2dztbHs07EenjyNUjD6d2DV2G-7l'
SCRIPT_DIR        = os.path.dirname(os.path.abspath(__file__))
GDRIVE_TOKEN      = os.path.join(SCRIPT_DIR, 'token.json')
GDRIVE_CREDS      = os.path.join(SCRIPT_DIR, 'credentials.json')
GDRIVE_SCOPES     = ['https://www.googleapis.com/auth/drive.file']

# ---------- Google Drive upload ----------

def upload_to_drive(local_path, filename):
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from google.auth.transport.requests import Request
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload

    creds = None
    if os.path.exists(GDRIVE_TOKEN):
        creds = Credentials.from_authorized_user_file(GDRIVE_TOKEN, GDRIVE_SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists(GDRIVE_CREDS):
                print(f"ERROR: credentials.json not found at {GDRIVE_CREDS}")
                print("Download OAuth2 credentials from Google Cloud Console and save there.")
                return
            flow = InstalledAppFlow.from_client_secrets_file(GDRIVE_CREDS, GDRIVE_SCOPES)
            creds = flow.run_local_server(port=0)
        with open(GDRIVE_TOKEN, 'w') as t:
            t.write(creds.to_json())

    service = build('drive', 'v3', credentials=creds)

    # Check if a file with this name already exists in the folder (update) or create new
    resp = service.files().list(
        q=f"name='{filename}' and '{GDRIVE_FOLDER_ID}' in parents and trashed=false",
        fields='files(id, name)'
    ).execute()
    existing = resp.get('files', [])

    mime  = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    media = MediaFileUpload(local_path, mimetype=mime, resumable=False)
    if existing:
        file_id = existing[0]['id']
        service.files().update(fileId=file_id, media_body=media).execute()
        print(f"Updated existing file on Drive: {filename} (id={file_id})")
    else:
        meta = {'name': filename, 'parents': [GDRIVE_FOLDER_ID]}
        f = service.files().create(body=meta, media_body=media, fields='id').execute()
        print(f"Uploaded new file to Drive: {filename} (id={f['id']})")

# ---------- normalisation ----------

def strip_accents(s):
    """Convert accented/diacritic characters to their ASCII base (à→a, é→e, ñ→n …)."""
    return ''.join(
        c for c in unicodedata.normalize('NFD', s)
        if unicodedata.category(c) != 'Mn'
    )

ADDR_EXPANSIONS = {
    r'\bst\b':    'street',      r'\bave\b':   'avenue',      r'\bblvd\b':  'boulevard',
    r'\bdr\b':    'drive',       r'\brd\b':    'road',        r'\bln\b':    'lane',
    r'\bpl\b':    'place',       r'\bct\b':    'court',       r'\bcir\b':   'circle',
    r'\bter\b':   'terrace',     r'\bpkwy\b':  'parkway',     r'\bhwy\b':   'highway',
    r'\bfwy\b':   'freeway',     r'\bsq\b':    'square',      r'\bste\b':   'suite',
    r'\bapt\b':   'apartment',   r'\bft\b':    'fort',        r'\bmt\b':    'mount',
    r'\bn\b':     'north',       r'\bs\b':     'south',
    r'\be\b':     'east',        r'\bw\b':     'west',
    r'\bne\b':    'northeast',   r'\bnw\b':    'northwest',
    r'\bse\b':    'southeast',   r'\bsw\b':    'southwest',
    r'\bjct\b':   'junction',    r'\bexpy\b':  'expressway',  r'\bexpwy\b': 'expressway',
    r'\bsaint\b': 'saint',
}

# Words stripped from address tokens before word-set comparison.
# Includes unit designators, articles, prepositions, AND expanded road-type words —
# road types ("street", "road", etc.) appear in almost every address so they must not
# be counted as meaningful overlap between two different streets.
ADDR_STOP = {
    'of', 'the', 'a', 'an', 'at', 'in', 'on', 'and',
    'suite', 'ste', 'apt', 'unit', 'floor',
    # expanded road-type words (expansion runs before core_words, so these are the full forms)
    'street', 'avenue', 'boulevard', 'drive', 'road', 'lane', 'place',
    'court', 'circle', 'terrace', 'parkway', 'highway', 'freeway',
    'square', 'expressway', 'junction', 'way', 'trail', 'trl',
}

NAME_EXPANSIONS = {
    r'&': 'and',
}

# Words that commonly appear in hotel names but carry no distinguishing meaning.
# Stripping these from both sides before comparing reveals the "core" brand/location tokens.
HOTEL_FILLER = {
    # property type
    'hotel', 'hotels', 'inn', 'inns', 'motel', 'motels', 'lodge', 'lodges',
    'resort', 'resorts', 'hostel', 'hostels', 'suites', 'suite', 'spa',
    'guesthouse', 'guest', 'house', 'rooms', 'room', 'bed', 'breakfast',
    # chain affiliation suffixes
    'by', 'an', 'a', 'the', 'and', 'at', 'in', 'of', 'on',
    # location/type qualifiers commonly added or dropped
    'airport', 'international', 'downtown', 'center', 'centre',
    'extended', 'stay', 'all', 'inclusive',
    # one-char tokens that slip through
}

def core_name_words(norm):
    """Return significant non-filler tokens from a pre-normalised name string."""
    return [w for w in norm.split() if w not in HOTEL_FILLER and len(w) > 1]

def name_match_status(bd_n, sd_n):
    """
    Specialised hotel-name comparison.
    GREEN  — core word sets have Jaccard ≥ 0.4, or the shorter set is a subset of the longer.
    YELLOW — at least one core word in common (partial / ambiguous).
    RED    — no core word in common at all.
    """
    bd_core = set(core_name_words(bd_n))
    sd_core = set(core_name_words(sd_n))

    if not bd_core and not sd_core: return 'green'
    if not bd_core or not sd_core:  return 'yellow'

    intersection = bd_core & sd_core
    union        = bd_core | sd_core

    # One name's core words are entirely contained in the other (e.g. "Microtel Dover"
    # ⊆ "Microtel Inn & Suites by Wyndham Dover") → same property, just different verbosity
    shorter = bd_core if len(bd_core) <= len(sd_core) else sd_core
    longer  = sd_core if len(bd_core) <= len(sd_core) else bd_core
    if shorter.issubset(longer):
        return 'green'

    jaccard = len(intersection) / len(union)
    if jaccard >= 0.4: return 'green'
    if jaccard >  0:   return 'yellow'
    return 'red'

# Common city abbreviations → full name
CITY_EXPANSIONS = {
    r'\bnyc\b':  'new york city', r'\bla\b':   'los angeles',
    r'\bsf\b':   'san francisco', r'\bchi\b':  'chicago',
    r'\bphx\b':  'phoenix',       r'\blas\b':  'las vegas',
    r'\bdc\b':   'washington',    r'\bdfw\b':  'dallas',
    r'\bmia\b':  'miami',         r'\bbos\b':  'boston',
    r'\bsea\b':  'seattle',       r'\batl\b':  'atlanta',
    r'\bden\b':  'denver',        r'\bhou\b':  'houston',
    r'\bord\b':  'chicago',       r'\bfll\b':  'fort lauderdale',
    r'\bmco\b':  'orlando',
}

def _expand(text, table):
    for pattern, replacement in table.items():
        text = re.sub(pattern, replacement, text)
    return text

def norm_addr(s):
    if not s: return ''
    s = strip_accents(str(s))
    s = s.lower()
    # Hyphens between digits become spaces (1-3-61 → 1 3 61); other hyphens → space too
    s = re.sub(r'-', ' ', s)
    s = re.sub(r'[^\w\s]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = _expand(s, ADDR_EXPANSIONS)
    return re.sub(r'\s+', ' ', s).strip()

def norm_name(s):
    if not s: return ''
    s = strip_accents(str(s))
    s = s.lower()
    s = re.sub(r"[''\-]", '', s)   # strip apostrophes/hyphens without adding space → O'Hare→ohare
    s = re.sub(r'[^\w\s&]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = _expand(s, NAME_EXPANSIONS)
    return re.sub(r'\s+', ' ', s).strip()

def norm_city(s):
    if not s: return ''
    s = strip_accents(str(s))
    s = s.lower()
    s = re.sub(r'[^\w\s]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = _expand(s, CITY_EXPANSIONS)
    return re.sub(r'\s+', ' ', s).strip()

def norm_phone(s):
    digits = re.sub(r'\D', '', str(s or ''))
    if len(digits) == 11 and digits.startswith('1'):
        digits = digits[1:]
    return digits

def norm_simple(s):
    if not s: return ''
    s = strip_accents(str(s))
    s = s.lower()
    s = re.sub(r'[^\w\s]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()

def similarity(a, b):
    return SequenceMatcher(None, a, b).ratio()

def word_overlap(a, b):
    """True if any significant word (>2 chars) from a appears in b, or any word if all are short."""
    wa = set(a.split())
    wb = set(b.split())
    sig_a = {w for w in wa if len(w) > 2}
    sig_b = {w for w in wb if len(w) > 2}
    if sig_a and sig_b:
        return bool(sig_a & sig_b)
    return bool(wa & wb)

def _words_prefix_match(w1, w2):
    """True if w1==w2, or one is a prefix of the other (min 4 chars — handles hill/hills)."""
    if w1 == w2: return True
    short, long = (w1, w2) if len(w1) < len(w2) else (w2, w1)
    return len(short) >= 4 and long.startswith(short)

def _core_addr_words(norm):
    """Significant address tokens: skip stop words and single characters."""
    return [w for w in norm.split() if w not in ADDR_STOP and len(w) > 1]

def _addr_subset(shorter, longer):
    """True if every word in `shorter` has a prefix-match partner in `longer`."""
    longer_list = list(longer)
    return all(any(_words_prefix_match(w, cand) for cand in longer_list) for w in shorter)

def _street_number(norm):
    """Return the leading street number token if the address starts with one, else None."""
    parts = norm.split()
    return parts[0] if parts and parts[0].isdigit() else None

def _meaningful_addr_words(core_words):
    """Non-numeric words from core — these identify the street name itself."""
    return {w for w in core_words if not w.isdigit()}

def _meaningful_overlap(bd_core, sd_core):
    """True if any non-numeric word from bd has a prefix-match partner in sd."""
    bd_m = _meaningful_addr_words(bd_core)
    sd_m = _meaningful_addr_words(sd_core)
    return any(_words_prefix_match(w1, w2) for w1 in bd_m for w2 in sd_m)

def addr_match_status(bd_n, sd_n):
    """
    GREEN  — core words of shorter address all found (prefix-matched) in longer,
             OR prefix-Jaccard >= 0.5, OR string similarity >= 0.80.
    YELLOW — street numbers differ but street name words overlap (nearby address),
             OR partial similarity.
    RED    — street numbers differ AND no street name word overlap,
             OR same/no numbers but no meaningful word overlap at all.
    """
    bd_num = _street_number(bd_n)
    sd_num = _street_number(sd_n)
    numbers_differ = bool(bd_num and sd_num and bd_num != sd_num)

    bd_core = _core_addr_words(bd_n)
    sd_core = _core_addr_words(sd_n)
    bd_set  = set(bd_core)
    sd_set  = set(sd_core)

    if not bd_set and not sd_set: return 'green'
    if not bd_set or not sd_set:  return 'yellow'

    if numbers_differ:
        # Different door numbers: yellow only when the street NAME words overlap
        # (e.g. 3535 vs 3545 Las Vegas Blvd — same street, plausible typo).
        # If even the street name differs → completely wrong address → red.
        return 'yellow' if _meaningful_overlap(bd_core, sd_core) else 'red'

    # Same (or absent) street number — check how well the rest matches.
    shorter, longer = (bd_core, sd_core) if len(bd_core) <= len(sd_core) else (sd_core, bd_core)
    if _addr_subset(set(shorter), set(longer)):
        return 'green'

    matched = sum(1 for w in bd_set if any(_words_prefix_match(w, c) for c in sd_set))
    if matched / len(bd_set | sd_set) >= 0.50:
        return 'green'

    sim = similarity(bd_n, sd_n)
    if sim >= 0.80: return 'green'

    # Final fallback: only non-numeric, non-road-type street name word overlap counts.
    # Character-level similarity is too blunt for addresses (e.g. "airport" vs "historic"
    # share scattered letters but are completely different streets).
    return 'yellow' if _meaningful_overlap(bd_core, sd_core) else 'red'

def compare(bd_val, sd_val, field='text'):
    """
    Returns 'green', 'yellow', or 'red'.
    RED only when values are completely different with zero word overlap.
    YELLOW when there is partial match or doubt.
    GREEN when values match well.
    """
    if field == 'phone':
        bd_n = norm_phone(bd_val)
        sd_n = norm_phone(sd_val)
        if not bd_n and not sd_n: return 'green'
        if not bd_n or not sd_n:  return 'yellow'
        if bd_n == sd_n:          return 'green'
        if bd_n in sd_n or sd_n in bd_n: return 'yellow'
        return 'red'

    if field == 'name':
        bd_n, sd_n = norm_name(bd_val), norm_name(sd_val)
        if not bd_n and not sd_n: return 'green'
        if not bd_n or not sd_n:  return 'yellow'
        if bd_n == sd_n:          return 'green'
        return name_match_status(bd_n, sd_n)

    if field == 'address':
        bd_n, sd_n = norm_addr(bd_val), norm_addr(sd_val)
        if not bd_n and not sd_n: return 'green'
        if not bd_n or not sd_n:  return 'yellow'
        if bd_n == sd_n:          return 'green'
        return addr_match_status(bd_n, sd_n)

    if field == 'city':
        bd_n, sd_n = norm_city(bd_val), norm_city(sd_val)
    else:
        bd_n, sd_n = norm_simple(bd_val), norm_simple(sd_val)

    if not bd_n and not sd_n: return 'green'
    if not bd_n or not sd_n:  return 'yellow'
    if bd_n == sd_n:          return 'green'

    sim = similarity(bd_n, sd_n)
    if sim >= 0.80: return 'green'
    if sim >= 0.50: return 'yellow'

    if word_overlap(bd_n, sd_n): return 'yellow'
    return 'red'

# ---------- mongo extractors ----------

def ati_info(doc):
    if not doc: return None
    a = doc.get('address', {})
    return dict(name=doc.get('vendorName',''), address=a.get('addressLine1',''),
                city=a.get('city',''), state=a.get('state',''),
                country=a.get('country',''), phone=a.get('phone',''))

def rapid_info(doc):
    if not doc: return None
    c = doc.get('contents', [])
    c = c[0] if isinstance(c, list) and c else (c or {})
    a = c.get('address', {})
    cn = a.get('CityName') or a.get('city','')
    city = cn.get('value','') if isinstance(cn, dict) else cn
    sp = a.get('StateProv') or a.get('state_province_code','')
    state = sp.get('StateCode','') if isinstance(sp, dict) else sp
    co = a.get('CountryName') or a.get('country_code','')
    country = co.get('Code','') if isinstance(co, dict) else co
    addr = a.get('line_1') or a.get('AddressLine1','')
    return dict(name=c.get('name',''), address=addr, city=city,
                state=state, country=country, phone=c.get('phone',''))

def sabre_info(doc):
    if not doc: return None
    loc  = doc.get('HotelDescriptiveInfo',{}).get('LocationInfo',{})
    addr = loc.get('Address',{})
    con  = loc.get('Contact',{})
    cn   = addr.get('CityName',{})
    sp   = addr.get('StateProv',{})
    co   = addr.get('CountryName',{})
    return dict(
        name    = doc.get('HotelInfo',{}).get('HotelName',''),
        address = addr.get('AddressLine1',''),
        city    = cn.get('value','') if isinstance(cn,dict) else cn,
        state   = sp.get('StateCode','') if isinstance(sp,dict) else sp,
        country = co.get('Code','') if isinstance(co,dict) else co,
        phone   = con.get('Phone','')
    )

# ---------- fetch data ----------

print(f"Connecting to MySQL…")
conn = pymysql.connect(host=MYSQL_HOST, port=MYSQL_PORT, user=MYSQL_USER,
                       password=MYSQL_PASSWORD, database=MYSQL_DB,
                       cursorclass=pymysql.cursors.DictCursor)
cur = conn.cursor()
cur.execute("""
    SELECT tb.id AS booking_id, tb.adapter_name, hb.provider_code,
           hb.hotel_id, hb.hotel_name, hb.hotel_address_line1, hb.hotel_address_city,
           hb.hotel_address_state_code, hb.hotel_address_country_code, hb.hotel_phone
    FROM TRAVEL_BOOKING tb
    JOIN HOTEL_BOOKING_V2 hb ON tb.id = hb.travel_booking_id
    WHERE DATE(tb.created_at) = %s
      AND tb.status IN ('Completed','Completed_CRS','Completed_Direct')
      AND tb.adapter_name IN ('SABRE_HOTEL','EAN_RAPID','ATI')
    ORDER BY tb.adapter_name, tb.id
""", (REPORT_DATE,))
bookings_raw = cur.fetchall()
conn.close()
print(f"Fetched {len(bookings_raw)} bookings for {REPORT_DATE}")

# Group by (adapter_name, provider_code) — same property booked multiple times becomes one row
from collections import defaultdict
_groups = defaultdict(list)
for b in bookings_raw:
    _groups[(b['adapter_name'], b['provider_code'] or '')].append(b)

bookings = []
for (adapter, code), group in _groups.items():
    rep = dict(group[0])
    rep['booking_ids'] = [str(g['booking_id']) for g in group]
    bookings.append(rep)

bookings.sort(key=lambda b: (b['adapter_name'], b['booking_ids'][0]))
print(f"Grouped into {len(bookings)} unique properties")

print("Connecting to MongoDB…")
mg = pymongo.MongoClient(MONGO_URI)['hotel']

ati_bk   = [b for b in bookings if b['adapter_name'] == 'ATI']
rapid_bk = [b for b in bookings if b['adapter_name'] == 'EAN_RAPID']
sabre_bk = [b for b in bookings if b['adapter_name'] == 'SABRE_HOTEL']

ati_props   = {p['vendorCode']: p for p in mg.atiProperty.find(
    {'vendorCode': {'$in': [b['provider_code'] for b in ati_bk]}})}

rapid_ids   = [int(b['provider_code']) for b in rapid_bk
               if b['provider_code'] and b['provider_code'].isdigit()]
rapid_props = {str(p['id']): p for p in mg.rapidProperty.find({'id': {'$in': rapid_ids}})}

sabre_props = {p['HotelInfo']['HotelCode']: p for p in mg.sabreProperty.find(
    {'HotelInfo.HotelCode': {'$in': [b['provider_code'] for b in sabre_bk]}})}

mg.client.close()

# ---------- build rows ----------

PROVIDER_LABEL = {'ATI': 'ATI', 'EAN_RAPID': 'Expedia/Rapid', 'SABRE_HOTEL': 'Sabre'}

def make_row(b):
    adapter = b['adapter_name']
    code    = b['provider_code'] or ''

    if adapter == 'ATI':
        supply = ati_info(ati_props.get(code))
    elif adapter == 'EAN_RAPID':
        supply = rapid_info(rapid_props.get(code))
    else:
        supply = sabre_info(sabre_props.get(code))

    bd = dict(
        name    = b['hotel_name'] or '',
        address = b['hotel_address_line1'] or '',
        city    = b['hotel_address_city'] or '',
        state   = b['hotel_address_state_code'] or '',
        country = b['hotel_address_country_code'] or '',
        phone   = b['hotel_phone'] or '',
    )

    if not supply:
        return dict(b=b, bd=bd, supply=None, statuses={}, row_status='yellow',
                    provider=PROVIDER_LABEL.get(adapter, adapter))

    statuses = {
        'name':    compare(bd['name'],    supply['name'],    field='name'),
        'address': compare(bd['address'], supply['address'], field='address'),
        'city':    compare(bd['city'],    supply['city'],    field='city'),
        'state':   compare(bd['state'],   supply['state'],   field='text'),
        'country': compare(bd['country'], supply['country'], field='text'),
        'phone':   compare(bd['phone'],   supply['phone'],   field='phone'),
    }
    row_status = 'red'    if 'red'    in statuses.values() else \
                 'yellow' if 'yellow' in statuses.values() else 'green'
    return dict(b=b, bd=bd, supply=supply, statuses=statuses, row_status=row_status,
                provider=PROVIDER_LABEL.get(adapter, adapter))


rows = [make_row(b) for b in bookings]

# Rows where every field matched (all-green) carry nothing to review — omit them
# from the report body. Summary counts below still reflect all rows evaluated.
display_rows = [r for r in rows if r['row_status'] != 'green']

# ---------- Excel ----------

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILL = {
    'green':  PatternFill('solid', fgColor='D4EDDA'),
    'yellow': PatternFill('solid', fgColor='FFF3CD'),
    'red':    PatternFill('solid', fgColor='F8D7DA'),
    'header': PatternFill('solid', fgColor='343A40'),
    'title':  PatternFill('solid', fgColor='1D3557'),
}
WHITE  = Font(color='FFFFFF', bold=True, size=11)
BOLD   = Font(bold=True)
WRAP   = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN   = Side(style='thin', color='CCCCCC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

summary       = {s: sum(1 for r in rows if r['row_status'] == s) for s in ('green','yellow','red')}
total_bookings = sum(len(r['b']['booking_ids']) for r in rows)

# Field pairs: (field_key, BD col index, SD col index)  — 1-based, starting after 4 fixed cols
FIELDS = ['name', 'address', 'city', 'state', 'country', 'phone']
FIELD_LABELS = ['Property Name', 'Address', 'City', 'State', 'Country', 'Phone']
# Fixed cols: Booking IDs(1) Provider(2) Provider Code(3) TST Hotel ID(4)
# Then pairs: BD Name(5) SD Name(6) | BD Addr(7) SD Addr(8) | ...
NCOLS = 4 + len(FIELDS) * 2

def write_excel(out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Property Report'

    # ---- Row 1: title ----
    ws.append([f'Hotel Booked Property Information — {DISPLAY_DATE}'])
    ws.merge_cells(f'A1:{get_column_letter(NCOLS)}1')
    c = ws['A1']
    c.fill  = FILL['title']
    c.font  = Font(color='FFFFFF', bold=True, size=14)
    c.alignment = CENTER
    ws.row_dimensions[1].height = 28

    # ---- Row 2: summary ----
    ws.append([
        f'Unique Properties: {len(rows)}   |   Total Bookings: {total_bookings}   |   '
        f'Rows shown below: {len(display_rows)} (all-match rows hidden)',
        '', '', '',
        f'✅ Match: {summary["green"]}', '',
        f'⚠️ Review: {summary["yellow"]}', '',
        f'❌ Wrong: {summary["red"]}',
    ])
    ws.merge_cells('A2:D2')
    ws.merge_cells('E2:F2')
    ws.merge_cells('G2:H2')
    ws['A2'].font = BOLD
    ws['E2'].fill = FILL['green'];  ws['E2'].font = BOLD; ws['E2'].alignment = CENTER
    ws['G2'].fill = FILL['yellow']; ws['G2'].font = BOLD; ws['G2'].alignment = CENTER
    ws['I2'].fill = FILL['red'];    ws['I2'].font = BOLD; ws['I2'].alignment = CENTER
    ws.row_dimensions[2].height = 20

    # ---- Row 3: legend ----
    ws.append([
        'BD = Booking Data (recorded at time of booking)',
        '', '', '',
        'SD = Supplier Data (from supplier reference database)',
    ])
    ws.merge_cells('A3:D3')
    ws.merge_cells('E3:J3')
    ws['A3'].font = Font(italic=True, size=10)
    ws['E3'].font = Font(italic=True, size=10)
    ws.row_dimensions[3].height = 16

    # ---- Row 4: spacer ----
    ws.append([])

    # ---- Row 5: column headers ----
    headers = ['Booking ID(s)', 'Provider', 'Provider Code', 'TST Hotel ID']
    for label in FIELD_LABELS:
        headers += [f'BD: {label}', f'SD: {label}']
    ws.append(headers)
    hrow = ws.max_row
    for col in range(1, NCOLS + 1):
        c = ws.cell(row=hrow, column=col)
        c.fill      = FILL['header']
        c.font      = WHITE
        c.alignment = CENTER
        c.border    = BORDER
    ws.row_dimensions[hrow].height = 36
    ws.freeze_panes = f'A{hrow + 1}'

    # ---- Data rows ----
    for r in display_rows:
        bd  = r['bd']
        sd  = r['supply']
        ids = ', '.join(r['b']['booking_ids'])
        hid = str(r['b'].get('hotel_id') or '')

        if not sd:
            row_vals = [ids, r['provider'], r['b']['provider_code'] or '', hid]
            for _ in FIELDS:
                row_vals += [bd.get(_, ''), '— not in Supplier Data —']
            ws.append(row_vals)
            drow = ws.max_row
            for col in range(1, NCOLS + 1):
                c = ws.cell(row=drow, column=col)
                c.fill   = FILL['yellow']
                c.border = BORDER
                c.alignment = WRAP
        else:
            s = r['statuses']
            row_vals = [ids, r['provider'], r['b']['provider_code'] or '', hid]
            for field in FIELDS:
                row_vals += [bd.get(field, ''), sd.get(field, '')]
            ws.append(row_vals)
            drow = ws.max_row
            # colour fixed cols neutral
            for col in range(1, 5):
                ws.cell(row=drow, column=col).border    = BORDER
                ws.cell(row=drow, column=col).alignment = WRAP
            # colour each BD/SD pair
            for i, field in enumerate(FIELDS):
                fill   = FILL[s[field]]
                bd_col = 5 + i * 2
                sd_col = bd_col + 1
                for col in (bd_col, sd_col):
                    c = ws.cell(row=drow, column=col)
                    c.fill      = fill
                    c.border    = BORDER
                    c.alignment = WRAP

    # ---- Column widths ----
    widths = [28, 14, 16, 14] + [28, 28] * len(FIELDS)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)
    print(f'Excel saved: {out_path}')

filename = f'Hotel_Booked_Property_Information_{DDMMYYYY}.xlsx'
out      = os.path.join(SCRIPT_DIR, filename)

write_excel(out)
upload_to_drive(out, filename)
